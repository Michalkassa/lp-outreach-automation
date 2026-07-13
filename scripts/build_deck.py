"""Build the pipeline presentation — Valori Capital.

16 slides, English, high-level: what it is, how it works, how to use it.
Design: Garamond, coral hero accent, sharp-edged cards, actor colours:
    NAVY = You (human)     CORAL = Claude (AI)     STEEL = files
Regenerate with:  .venv/bin/python scripts/build_deck.py
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ── Palette ──────────────────────────────────────────────────────────
CORAL     = RGBColor(0xEF, 0x7B, 0x88)   # Claude / AI
CORAL_LT  = RGBColor(0xF4, 0xA5, 0xAE)
CORAL_XLT = RGBColor(0xFB, 0xE4, 0xE7)
NAVY      = RGBColor(0x2E, 0x3A, 0x50)   # You
NAVY_XLT  = RGBColor(0xE9, 0xEC, 0xF2)
STEEL     = RGBColor(0x7E, 0x92, 0xA9)   # files
STEEL_LT  = RGBColor(0xB4, 0xC0, 0xCF)
STEEL_XLT = RGBColor(0xF0, 0xF3, 0xF7)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF8, 0xF8, 0xF8)
LIGHT_GR  = RGBColor(0xE8, 0xE8, 0xE8)
MID_GRAY  = RGBColor(0x8A, 0x8A, 0x8A)
DARK_GRAY = RGBColor(0x2C, 0x2C, 0x2C)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

SERIF = "Garamond"
SANS  = "Garamond"
MONO  = "Consolas"

I = Inches

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ── Helpers ──────────────────────────────────────────────────────────
def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _track(par, spc):
    for r in par.runs:
        r.font._rPr.set("spc", str(spc))


def rect(slide, l, t, w, h, fill=CORAL, border=None, bw=Pt(1)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if border:
        sh.line.color.rgb = border
        sh.line.width = bw
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def box(slide, l, t, w, h, fill=WHITE, border=None, bw=Pt(1)):
    """Sharp-edged box (all presentation edges are square)."""
    return rect(slide, l, t, w, h, fill=fill, border=border, bw=bw)


def shape(slide, kind, l, t, w, h, fill=CORAL, border=None, bw=Pt(1), rot=0):
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if border:
        sh.line.color.rgb = border
        sh.line.width = bw
    else:
        sh.line.fill.background()
    if rot:
        sh.rotation = rot
    sh.shadow.inherit = False
    return sh


def txt(slide, l, t, w, h, text, sz=18, color=DARK_GRAY, bold=False,
        align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP,
        italic=False, tracking=None, line_spacing=None):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    if tracking:
        _track(p, tracking)
    return tb


def multi(slide, l, t, w, h, lines, sz=12, color=DARK_GRAY, bold=False,
          align=PP_ALIGN.LEFT, font=SANS, spacing=1.5, italic=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(sz)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.italic = italic
        p.font.name = font
        p.alignment = align
        p.space_after = Pt(sz * (spacing - 1))
    return tb


def boxtext(sh, text, sz=12, color=WHITE, bold=False, font=SANS,
            align=PP_ALIGN.CENTER, sub=None, sub_sz=9, sub_color=None,
            sub_font=None):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    if sub is not None:
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(sub_sz)
        p2.font.color.rgb = sub_color or color
        p2.font.name = sub_font or MONO
        p2.alignment = align
    return sh


def add_notes(slide, lines):
    slide.notes_slide.notes_text_frame.text = "\n".join(
        "• " + ln for ln in lines)


def _head(slide, x, y, size, color, rot):
    shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE,
          x - size / 2, y - size / 2, size, size, fill=color, rot=rot)


def harrow(slide, x1, x2, y, color=MID_GRAY, thick=Pt(1.6), head=I(0.14)):
    if x2 >= x1:
        rect(slide, x1, y - thick / 2, x2 - x1 - head / 2, thick, fill=color)
        _head(slide, x2 - head / 2, y, head, color, 90)
    else:
        rect(slide, x2 + head / 2, y - thick / 2, x1 - x2 - head / 2, thick,
             fill=color)
        _head(slide, x2 + head / 2, y, head, color, 270)


def varrow(slide, y1, y2, x, color=MID_GRAY, thick=Pt(1.6), head=I(0.14)):
    if y2 >= y1:
        rect(slide, x - thick / 2, y1, thick, y2 - y1 - head / 2, fill=color)
        _head(slide, x, y2 - head / 2, head, color, 180)
    else:
        rect(slide, x - thick / 2, y2 + head / 2, thick, y1 - y2 - head / 2,
             fill=color)
        _head(slide, x, y2 + head / 2, head, color, 0)


def hline(slide, x1, x2, y, color=MID_GRAY, thick=Pt(1.6)):
    rect(slide, min(x1, x2), y - thick / 2, abs(x2 - x1), thick, fill=color)


def vline(slide, y1, y2, x, color=MID_GRAY, thick=Pt(1.6)):
    rect(slide, x - thick / 2, min(y1, y2), thick, abs(y2 - y1), fill=color)


def icon_human(slide, x, y, size=I(0.3)):
    shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill=NAVY)


def icon_ai(slide, x, y, size=I(0.3)):
    shape(slide, MSO_SHAPE.DIAMOND, x, y, size, size, fill=CORAL)


def icon_script(slide, x, y, size=I(0.3)):
    rect(slide, x, y, size, size, fill=STEEL)


def actor_legend(slide, x, y, sz=9.5, gap=I(1.8)):
    items = [(icon_human, "You"), (icon_ai, "Claude (AI)"),
             (icon_script, "Files")]
    for i, (fn, label) in enumerate(items):
        ix = x + i * gap
        fn(slide, ix, y, I(0.16))
        txt(slide, ix + I(0.24), y - I(0.055), I(1.0), I(0.28), label,
            sz=sz, color=MID_GRAY)


def footer(slide):
    txt(slide, I(8.5), I(7.13), I(4.6), I(0.25),
        "VALORI CAPITAL · LP OUTREACH", sz=7.5, color=MID_GRAY,
        align=PP_ALIGN.RIGHT, tracking=180)


def plain_slide(kicker, title):
    s = prs.slides.add_slide(BLANK)
    set_bg(s, WHITE)
    txt(s, I(0.9), I(0.42), I(11.8), I(0.3), kicker.upper(), sz=11,
        color=CORAL, bold=True, tracking=220)
    txt(s, I(0.9), I(0.72), I(11.8), I(0.6), title, sz=26, color=DARK_GRAY,
        bold=True, font=SERIF)
    rect(s, I(0.92), I(1.42), I(0.9), Pt(2.5), fill=CORAL)
    footer(s)
    return s


def shot(slide, l, t, w, h, code, spec_lines):
    sh = box(slide, l, t, w, h, fill=OFF_WHITE, border=MID_GRAY, bw=Pt(1))
    sh.line.dash_style = MSO_LINE.DASH
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "✂ SCREENSHOT %s" % code
    p.font.size = Pt(13)
    p.font.color.rgb = CORAL
    p.font.bold = True
    p.font.name = SANS
    p.alignment = PP_ALIGN.CENTER
    for line in spec_lines:
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.size = Pt(10)
        p2.font.color.rgb = MID_GRAY
        p2.font.italic = True
        p2.font.name = SANS
        p2.alignment = PP_ALIGN.CENTER
    return sh


def chip(slide, l, t, w, h, text, fill=OFF_WHITE, color=DARK_GRAY,
         border=LIGHT_GR, sz=10, bold=False, font=SANS):
    sh = box(slide, l, t, w, h, fill=fill, border=border, bw=Pt(1))
    boxtext(sh, text, sz=sz, color=color, bold=bold, font=font)
    return sh


def card(slide, l, t, w, h, accent=CORAL):
    c = box(slide, l, t, w, h, fill=WHITE, border=LIGHT_GR, bw=Pt(1))
    rect(slide, l, t, w, Pt(3), fill=accent)
    return c


def node(slide, l, t, w, h, label, sub=None, actor=CORAL, sz=10, sub_sz=8):
    n = box(slide, l, t, w, h, fill=WHITE, border=LIGHT_GR, bw=Pt(1))
    rect(slide, l, t, I(0.07), h, fill=actor)
    boxtext(n, label, sz=sz, color=DARK_GRAY, bold=True,
            sub=sub, sub_sz=sub_sz, sub_color=MID_GRAY)
    return n


def file_node(slide, l, t, w, h, label, sub=None, sz=9):
    n = box(slide, l, t, w, h, fill=STEEL_XLT, border=STEEL_LT, bw=Pt(1))
    boxtext(n, label, sz=sz, color=NAVY, bold=False, font=MONO,
            sub=sub, sub_sz=7.5, sub_color=STEEL)
    return n


def diamond(slide, cx, cy, size, label, sz=8.5):
    d = shape(slide, MSO_SHAPE.DIAMOND, cx - size / 2, cy - size / 2,
              size, size, fill=WHITE, border=NAVY, bw=Pt(1.4))
    boxtext(d, label, sz=sz, color=NAVY, bold=True)
    return d


# ════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s, CORAL)
shape(s, MSO_SHAPE.OVAL, I(8.6), I(-1.6), I(7.4), I(7.4),
      fill=None, border=CORAL_LT, bw=Pt(2))
shape(s, MSO_SHAPE.OVAL, I(9.6), I(-0.6), I(5.4), I(5.4),
      fill=None, border=CORAL_LT, bw=Pt(1.25))
txt(s, I(1.1), I(1.9), I(8), I(0.35), "VALORI CAPITAL", sz=13, color=WHITE,
    bold=True, tracking=300)
multi(s, I(1.05), I(2.35), I(10.5), I(2.2),
      ["Semi-Automatic", "Fundraising Pipeline"],
      sz=46, color=WHITE, bold=True, font=SERIF, spacing=1.0)
rect(s, I(1.12), I(4.45), I(1.2), Pt(2.5), fill=WHITE)
txt(s, I(1.1), I(4.75), I(9.5), I(0.5),
    "The machine prepares. You decide every send.", sz=16, color=WHITE)
add_notes(s, ["One sentence: AI does the legwork, a person sends."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 2 — What it is
# ════════════════════════════════════════════════════════════════════
s = plain_slide("What it is", "Three actors, one rule")
actors = [
    (icon_human, NAVY, "You", [
        "Decide who to contact",
        "Edit every draft",
        "Send every email"]),
    (icon_ai, CORAL, "Claude (AI)", [
        "Researches each investor",
        "Scores the fit",
        "Writes the drafts"]),
    (icon_script, STEEL, "Files", [
        "Everything lives in one folder",
        "Nothing is hidden",
        "Nothing gets lost"]),
]
for i, (icon_fn, color, name, lines) in enumerate(actors):
    cx = I(0.9) + i * I(4.0)
    card(s, cx, I(1.95), I(3.7), I(3.15), accent=color)
    icon_fn(s, cx + I(0.35), I(2.35), I(0.5))
    txt(s, cx + I(1.05), I(2.36), I(2.5), I(0.5), name, sz=19,
        color=DARK_GRAY, bold=True, font=SERIF)
    multi(s, cx + I(0.38), I(3.15), I(3.05), I(1.8),
          ["–  " + ln for ln in lines], sz=11.5, spacing=1.6)
banner = box(s, I(0.9), I(5.5), I(11.5), I(0.85), fill=NAVY)
boxtext(banner,
        "The AI never touches a mailbox — every email is sent by "
        "a person.", sz=15, color=WHITE, bold=True)
add_notes(s, ["Set the colour code: navy you, coral AI, grey files."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 3 — How it works (the loop)
# ════════════════════════════════════════════════════════════════════
s = plain_slide("How it works", "Five steps, one loop")
steps = [
    ("1", "Add prospects", "one spreadsheet", NAVY),
    ("2", "Research & score", "/research", CORAL),
    ("3", "Draft", "/draft", CORAL),
    ("4", "Review & send", "your mailbox", NAVY),
    ("5", "Learn", "/calibrate", CORAL),
]
BOX_W, BOX_H = I(3.1), I(1.25)
TOP_Y, BOT_Y = I(2.1), I(4.75)
top_xs = [I(1.3), I(5.1), I(8.9)]        # steps 1,2,3
bot_xs = [I(6.95), I(3.15)]              # steps 4,5 (right to left)

def loop_node(x, y, step):
    num, label, cmd, actor = steps[step]
    box(s, x, y, BOX_W, BOX_H, fill=WHITE, border=LIGHT_GR, bw=Pt(1))
    rect(s, x, y, I(0.08), BOX_H, fill=actor)
    txt(s, x + I(0.24), y + I(0.12), I(0.7), I(0.5), num, sz=24,
        color=actor, font=SERIF)
    txt(s, x + I(0.78), y + I(0.26), BOX_W - I(0.95), I(0.35), label,
        sz=14.5, color=DARK_GRAY, bold=True)
    txt(s, x + I(0.78), y + I(0.7), BOX_W - I(0.95), I(0.3), cmd,
        sz=9.5, color=MID_GRAY, font=MONO)

for i, x in enumerate(top_xs):
    loop_node(x, TOP_Y, i)
    if i < 2:
        harrow(s, x + BOX_W, top_xs[i + 1], TOP_Y + BOX_H / 2, color=CORAL)
for i, x in enumerate(bot_xs):
    loop_node(x, BOT_Y, 3 + i)
harrow(s, bot_xs[0], bot_xs[1] + BOX_W, BOT_Y + BOX_H / 2, color=CORAL)
varrow(s, TOP_Y + BOX_H, BOT_Y + BOX_H / 2, top_xs[2] + BOX_W / 2,
       color=CORAL)
hline(s, bot_xs[0] + BOX_W, top_xs[2] + BOX_W / 2, BOT_Y + BOX_H / 2,
      color=CORAL, thick=Pt(1.6))
hline(s, top_xs[0] + BOX_W / 2, bot_xs[1], BOT_Y + BOX_H / 2,
      color=CORAL, thick=Pt(1.6))
varrow(s, BOT_Y + BOX_H / 2, TOP_Y + BOX_H, top_xs[0] + BOX_W / 2,
       color=CORAL)
txt(s, I(3.9), I(3.6), I(5.5), I(0.35),
    "every round, the drafts need fewer of your edits",
    sz=11.5, color=MID_GRAY, align=PP_ALIGN.CENTER, italic=True)
actor_legend(s, I(4.6), I(6.55))
add_notes(s, ["Walk the loop once — the rest of the deck is one slide "
              "per step."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 4 — Step 1: add prospects
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 1 · one spreadsheet", "Add prospects")
shot(s, I(0.9), I(1.75), I(6.9), I(4.6), "4-A", [
    "the prospect sheet in Excel",
    "columns: company, country, contact",
    "a few rows, anonymised",
])
points = [
    "One row per investor",
    "Name, country, contact — that is all",
    "Claude sets up the rest by itself",
]
for i, ptxt in enumerate(points):
    cy = I(2.15) + i * I(1.25)
    card(s, I(8.2), cy, I(4.25), I(1.0), accent=NAVY)
    txt(s, I(8.5), cy + I(0.3), I(3.7), I(0.5), ptxt, sz=12.5,
        color=DARK_GRAY, bold=True)
add_notes(s, ["The only data entry in the whole system."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 5 — Step 1 continued: import
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 1 · one command", "Pull the prospects in")
cmd = box(s, I(0.9), I(1.8), I(4.6), I(0.75), fill=DARK_GRAY)
boxtext(cmd, "/import-lps", sz=15, color=WHITE, font=MONO)
imp = [
    "Reads the sheet, creates one file per investor",
    "Skips rows already imported — safe to re-run",
    "No research yet — that is step 2",
]
for i, o in enumerate(imp):
    ry = I(3.0) + i * I(0.78)
    tick = shape(s, MSO_SHAPE.OVAL, I(1.0), ry + I(0.04), I(0.3), I(0.3),
                 fill=CORAL)
    boxtext(tick, "✓", sz=11, color=WHITE, bold=True)
    txt(s, I(1.5), ry + I(0.02), I(4.1), I(0.6), o, sz=12.5,
        color=DARK_GRAY, bold=True)
card(s, I(0.9), I(5.55), I(4.6), I(1.15), accent=NAVY)
txt(s, I(1.2), I(5.72), I(4.05), I(0.35), "Adding just one investor?",
    sz=12.5, color=DARK_GRAY, bold=True)
txt(s, I(1.2), I(6.12), I(4.05), I(0.5),
    "Skip the sheet — type /research <name> directly.", sz=10.5,
    color=MID_GRAY)
shot(s, I(6.1), I(1.8), I(6.35), I(4.9), "5-A", [
    "Claude right after /import-lps",
    "the summary list of imported investors",
    "anonymise company names",
])
add_notes(s, ["One command turns the sheet into the working pipeline.",
              "Single investors can go straight to /research."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 6 — Step 2: research
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 2 · one command", "Claude researches every investor")
cmd = box(s, I(0.9), I(1.8), I(4.6), I(0.75), fill=DARK_GRAY)
boxtext(cmd, "/research  <investor name>", sz=15, color=WHITE, font=MONO)
out = [
    "Who they are and what they invest in",
    "The right person to write to",
    "3 personal angles for the email",
    "A fit score from 1 to 100",
]
for i, o in enumerate(out):
    ry = I(3.0) + i * I(0.78)
    tick = shape(s, MSO_SHAPE.OVAL, I(1.0), ry + I(0.04), I(0.3), I(0.3),
                 fill=CORAL)
    boxtext(tick, "✓", sz=11, color=WHITE, bold=True)
    txt(s, I(1.5), ry + I(0.02), I(4.1), I(0.6), o, sz=12.5,
        color=DARK_GRAY, bold=True)
txt(s, I(0.9), I(6.3), I(4.6), I(0.4),
    "A few minutes per investor, every fact sourced.", sz=10.5,
    color=MID_GRAY, italic=True)
shot(s, I(6.1), I(1.8), I(6.35), I(4.9), "6-A", [
    "Claude running /research",
    "the finished investor file next to it",
    "anonymise the name",
])
add_notes(s, ["You type one line; the dossier appears as a file."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 7 — The score
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Simple priorities", "The score tells you where to focus")
tiers = [
    ("75+", "Draft now", CORAL, WHITE),
    ("35 – 74", "Check the gaps, then draft", CORAL_LT, WHITE),
    ("below 35", "Leave for later", OFF_WHITE, MID_GRAY),
]
for i, (rng, label, bg, fg) in enumerate(tiers):
    ry = I(2.1) + i * I(1.15)
    box(s, I(0.9), ry, I(5.3), I(0.95), fill=bg)
    txt(s, I(1.2), ry + I(0.22), I(1.7), I(0.5), rng, sz=19, color=fg,
        bold=True, font=SERIF)
    txt(s, I(3.0), ry + I(0.3), I(3.1), I(0.4), label, sz=13, color=fg,
        bold=True)
txt(s, I(0.9), I(5.75), I(5.3), I(0.6),
    "One glance shows what to work on today.", sz=10.5,
    color=MID_GRAY, italic=True)
shot(s, I(6.9), I(1.8), I(5.55), I(4.9), "7-A", [
    "one investor file with its",
    "score at the top",
    "anonymise the name",
])
add_notes(s, ["Scores rank the pipeline; the human picks from the top."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 8 — output.xlsx: the pipeline at a glance
# ════════════════════════════════════════════════════════════════════
s = plain_slide("One file to watch", "output.xlsx — your pipeline "
                "at a glance")
points7 = [
    ("Every investor on one line", "sorted by score, best on top"),
    ("Colour shows the stage", "researched · drafted · sent"),
    ("Always up to date", "refreshes after every step"),
]
for i, (head, sub) in enumerate(points7):
    cy = I(2.0) + i * I(1.45)
    card(s, I(0.9), cy, I(4.5), I(1.25), accent=STEEL)
    txt(s, I(1.2), cy + I(0.2), I(3.95), I(0.35), head, sz=13,
        color=DARK_GRAY, bold=True)
    txt(s, I(1.2), cy + I(0.62), I(3.95), I(0.35), sub, sz=10.5,
        color=MID_GRAY)
txt(s, I(0.9), I(6.45), I(4.5), I(0.4),
    "Open it like any Excel file.", sz=10.5, color=MID_GRAY, italic=True)
shot(s, I(6.1), I(1.8), I(6.35), I(4.9), "8-A", [
    "output.xlsx in Excel",
    "colour-coded rows, scores, stages",
    "anonymise company names",
])
add_notes(s, ["The management view: whole pipeline in one spreadsheet."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 9 — Step 3: draft (and how drafting works)
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 3 · one command",
                "Ready emails per investor — as many as you want")
cmd = box(s, I(0.9), I(1.8), I(4.6), I(0.95), fill=DARK_GRAY)
boxtext(cmd, "/draft  <investor name>", sz=15, color=WHITE, font=MONO,
        sub="want fewer than 3?   /draft <name> 2", sub_sz=9.5,
        sub_color=LIGHT_GR)
card(s, I(0.9), I(2.95), I(4.6), I(2.1), accent=CORAL)
txt(s, I(1.2), I(3.15), I(4.0), I(0.35), "How a draft is built", sz=13.5,
    color=DARK_GRAY, bold=True)
multi(s, I(1.2), I(3.6), I(4.05), I(1.3), [
    "–  Opens with one true fact about them",
    "–  The rest is the same proven structure",
    "–  One angle per draft — 1, 2 or 3 (A, B, C)",
], sz=11, spacing=1.5)
multi(s, I(0.9), I(5.25), I(4.6), I(1.2), [
    "Right language automatically.",
    "You pick the best one — or merge two.",
], sz=11, color=MID_GRAY, spacing=1.4, italic=True)
shot(s, I(6.1), I(1.8), I(6.35), I(2.6), "9-A", [
    "two draft variants side by side",
    "anonymise the recipient",
])
shot(s, I(6.1), I(4.6), I(6.35), I(2.1), "9-B", [
    "one email with the personal opening highlighted",
])
add_notes(s, ["Personalisation is one true fact — that is what makes "
              "it land."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 10 — The review packet
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Before you send", "The review packet — every draft "
                "on one desk")
cmd = box(s, I(0.9), I(1.8), I(4.6), I(0.75), fill=DARK_GRAY)
boxtext(cmd, "/review-packet", sz=15, color=WHITE, font=MONO)
points9 = [
    ("One file, every pending email", "nothing to hunt for"),
    ("Summary table on top", "who, angle, length"),
    ("Read on screen or print", "made for a review session"),
]
for i, (head, sub) in enumerate(points9):
    cy = I(3.0) + i * I(1.25)
    card(s, I(0.9), cy, I(4.6), I(1.05), accent=CORAL)
    txt(s, I(1.2), cy + I(0.15), I(4.05), I(0.35), head, sz=12.5,
        color=DARK_GRAY, bold=True)
    txt(s, I(1.2), cy + I(0.55), I(4.05), I(0.35), sub, sz=10.5,
        color=MID_GRAY)
shot(s, I(6.1), I(1.8), I(6.35), I(4.9), "10-A", [
    "the review packet",
    "summary table on top, first email below",
    "anonymise names",
])
add_notes(s, ["Compare all pending emails side by side in one sitting."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 11 — Step 4: read, edit, send
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 4 · you decide", "Read, edit, send")
flow = [
    ("You read and pick", "edit anything", NAVY),
    ("You send it", "from your own mailbox", NAVY),
    ("You paste back what you sent", "/log-sent", NAVY),
]
for i, (label, sub, actor) in enumerate(flow):
    x = I(0.9) + i * I(4.1)
    node(s, x, I(2.0), I(3.6), I(1.1), label, sub=sub, actor=actor,
         sz=13, sub_sz=10)
    if i < 2:
        harrow(s, x + I(3.6), x + I(4.1), I(2.55), color=NAVY)
banner = box(s, I(0.9), I(3.7), I(11.5), I(0.85), fill=NAVY)
boxtext(banner, "The AI never sends — every email leaves your mailbox.",
        sz=15, color=WHITE, bold=True)
shot(s, I(0.9), I(4.95), I(11.5), I(1.85), "11-A", [
    "your mail client with the finished email pasted in, deck attached — "
    "anonymise the recipient",
])
add_notes(s, ["The one non-negotiable rule, stated plainly.",
              "Pasting back what you sent feeds the learning step."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 12 — Step 5: it learns
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Step 5 · it learns", "Your edits teach it")
learn = [
    ("It compares", "your sent version vs its draft", CORAL),
    ("It spots patterns", "what you always change", CORAL),
    ("Next drafts follow your style", "automatically", CORAL),
]
for i, (label, sub, actor) in enumerate(learn):
    x = I(0.9) + i * I(4.1)
    node(s, x, I(1.9), I(3.6), I(1.1), label, sub=sub, actor=actor,
         sz=13, sub_sz=10)
    if i < 2:
        harrow(s, x + I(3.6), x + I(4.1), I(2.45), color=CORAL)
ladder = [
    ("seen once", "noticed", OFF_WHITE, MID_GRAY),
    ("seen twice", "becomes a rule", CORAL_LT, WHITE),
    ("seen three times", "it asks you to update the style", CORAL, WHITE),
]
for i, (freq, result, bg, fg) in enumerate(ladder):
    cx = I(0.9) + i * I(4.1)
    b = box(s, cx, I(3.5), I(3.6), I(0.85), fill=bg)
    boxtext(b, freq, sz=11, color=fg, bold=True, sub=result, sub_sz=9.5,
            sub_color=fg, sub_font=SANS)
shot(s, I(0.9), I(4.85), I(11.5), I(1.9), "12-A", [
    "a before / after: Claude's draft next to the version you sent, "
    "one repeated edit visible",
])
add_notes(s, ["The longer you use it, the less you correct.",
              "Style changes are proposed, never made silently."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 13 — Weekly review
# ════════════════════════════════════════════════════════════════════
s = plain_slide("Once a week", "/weekly-review — one command, "
                "one status page")
wr = [
    ("What moved", "sent, replied, waiting"),
    ("What works", "which opening angle gets replies"),
    ("What needs you", "replies awaiting an answer"),
    ("What went stale", "research worth refreshing"),
]
for i, (head, sub) in enumerate(wr):
    cx = I(0.9) + (i % 2) * I(2.95)
    cy = I(2.0) + (i // 2) * I(1.5)
    card(s, cx, cy, I(2.75), I(1.3), accent=CORAL)
    txt(s, cx + I(0.22), cy + I(0.18), I(2.35), I(0.35), head, sz=12.5,
        color=DARK_GRAY, bold=True)
    txt(s, cx + I(0.22), cy + I(0.62), I(2.35), I(0.55), sub, sz=10,
        color=MID_GRAY)
txt(s, I(0.9), I(5.25), I(5.7), I(0.9),
    "Five minutes on a Monday — then you know exactly "
    "which investors to add next.", sz=11.5, color=MID_GRAY,
    italic=True, line_spacing=1.3)
shot(s, I(7.0), I(1.8), I(5.45), I(4.9), "13-A", [
    "the weekly review output",
    "counts per stage and reply rates visible",
    "anonymise names",
])
add_notes(s, ["Weekly rhythm: read the status, refill the list."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 14 — The commands
# ════════════════════════════════════════════════════════════════════
s = plain_slide("The entire interface", "A handful of commands — that is all")
cmds = [
    ("/import-lps", "pull new prospects in from the sheet"),
    ("/research <name>", "research and score one investor"),
    ("/draft <name>", "write the drafts — 1, 2 or 3 variants"),
    ("/review-packet", "collect all drafts into one file"),
    ("/log-sent <name>", "store what you actually sent"),
    ("/calibrate", "learn from your edits"),
    ("/weekly-review", "a one-page status of the pipeline"),
]
for i, (cmd, desc) in enumerate(cmds):
    ry = I(1.95) + i * I(0.68)
    if i % 2 == 0:
        rect(s, I(0.9), ry - I(0.05), I(11.5), I(0.62), fill=OFF_WHITE)
    txt(s, I(1.15), ry + I(0.09), I(3.1), I(0.35), cmd, sz=13.5,
        color=DARK_GRAY, bold=True, font=MONO)
    txt(s, I(4.6), ry + I(0.11), I(7.5), I(0.35), desc, sz=12.5,
        color=MID_GRAY)
add_notes(s, ["Plain words typed into a chat — no software to learn."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 15 — The whole system on one page
# ════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s, WHITE)
txt(s, I(0.9), I(0.3), I(8), I(0.28), "THE BIG PICTURE", sz=11,
    color=CORAL, bold=True, tracking=220)
txt(s, I(0.9), I(0.6), I(11.5), I(0.5),
    "The whole system on one page", sz=24, color=DARK_GRAY, bold=True,
    font=SERIF)
footer(s)

LANE_X, LANE_W = I(0.82), I(12.42)
Y_Y, Y_H = I(1.45), I(1.7)      # you
C_Y, C_H = I(3.25), I(1.7)      # claude
F_Y, F_H = I(5.05), I(1.15)     # files
rect(s, LANE_X, Y_Y, LANE_W, Y_H, fill=NAVY_XLT)
rect(s, LANE_X, C_Y, LANE_W, C_H, fill=CORAL_XLT)
rect(s, LANE_X, F_Y, LANE_W, F_H, fill=STEEL_XLT)
for label, ly, lh, color in [("YOU", Y_Y, Y_H, NAVY),
                             ("CLAUDE", C_Y, C_H, CORAL),
                             ("FILES", F_Y, F_H, STEEL)]:
    tb = txt(s, I(0.02), ly + lh / 2 - I(0.6), I(0.75), I(1.2), label,
             sz=9, color=color, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, tracking=160)
    tb.rotation = 270

NW, NH = I(2.0), I(0.72)
y_n = Y_Y + I(0.49)
c_n = C_Y + I(0.49)
f_n = F_Y + I(0.28)

# you lane
node(s, I(1.1), y_n, NW, NH, "Add prospects", sub="one sheet",
     actor=NAVY, sz=10.5)
node(s, I(7.0), y_n, NW, NH, "Read + pick", sub="edit anything",
     actor=NAVY, sz=10.5)
d_ok = diamond(s, I(9.65), y_n + NH / 2, I(0.8), "OK?", sz=8.5)
node(s, I(10.5), y_n, NW, NH, "Send it", sub="your mailbox",
     actor=NAVY, sz=10.5)
harrow(s, I(9.0), I(9.25), y_n + NH / 2, color=NAVY, thick=Pt(1.3))
harrow(s, I(10.05), I(10.5), y_n + NH / 2, color=NAVY, thick=Pt(1.3))
txt(s, I(10.06), y_n - I(0.24), I(0.5), I(0.22), "yes", sz=8,
    color=MID_GRAY, italic=True)
vline(s, y_n - I(0.16), y_n + I(0.06), I(9.65), color=MID_GRAY, thick=Pt(1))
hline(s, I(8.0), I(9.65), y_n - I(0.16), color=MID_GRAY, thick=Pt(1))
varrow(s, y_n - I(0.16), y_n, I(8.0), color=MID_GRAY, thick=Pt(1))
txt(s, I(8.35), y_n - I(0.42), I(1.3), I(0.22), "no: edit", sz=8,
    color=MID_GRAY, italic=True)

# claude lane
node(s, I(1.1), c_n, NW, NH, "Research + score", sub="/research",
     actor=CORAL, sz=10.5)
node(s, I(3.75), c_n, NW, NH, "Draft ×1–3", sub="/draft",
     actor=CORAL, sz=10.5)
node(s, I(6.4), c_n, NW, NH, "Collect drafts", sub="/review-packet",
     actor=CORAL, sz=10.5)
node(s, I(10.5), c_n, NW, NH, "Learn your edits", sub="/calibrate",
     actor=CORAL, sz=10.5)
harrow(s, I(3.1), I(3.75), c_n + NH / 2, color=CORAL, thick=Pt(1.3))
harrow(s, I(5.75), I(6.4), c_n + NH / 2, color=CORAL, thick=Pt(1.3))

# cross-lane
varrow(s, y_n + NH, c_n, I(2.1), color=NAVY, thick=Pt(1.3))
varrow(s, c_n, y_n + NH, I(7.4), color=CORAL, thick=Pt(1.3))
varrow(s, y_n + NH, c_n, I(11.5), color=NAVY, thick=Pt(1.3))

# feedback: learn -> draft
fb_y = C_Y + C_H - I(0.14)
vline(s, c_n + NH, fb_y, I(11.9), color=CORAL, thick=Pt(1.2))
hline(s, I(4.75), I(11.9), fb_y, color=CORAL, thick=Pt(1.2))
varrow(s, fb_y, c_n + NH, I(4.75), color=CORAL, thick=Pt(1.2))
txt(s, I(7.3), fb_y - I(0.26), I(3.6), I(0.22),
    "next drafts follow your style", sz=8.5, color=CORAL, italic=True)

# files lane
files = [
    (I(1.1), "prospect sheet"),
    (I(3.75), "investor files + scores"),
    (I(6.4), "drafts"),
    (I(10.5), "sent archive"),
]
for x, label in files:
    file_node(s, x, f_n, NW, I(0.6), label, sz=9)

txt(s, I(0.9), I(6.5), I(11.5), I(0.35),
    "Everything is a file in one folder — open, readable, yours.",
    sz=10.5, color=MID_GRAY, align=PP_ALIGN.CENTER, italic=True)
actor_legend(s, I(4.4), I(7.12), sz=8, gap=I(1.5))
add_notes(s, ["One picture: you top, AI middle, files below.",
              "Point at the feedback arrow — it improves on its own."])

# ════════════════════════════════════════════════════════════════════
# SLIDE 16 — Closing
# ════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s, CORAL)
shape(s, MSO_SHAPE.OVAL, I(-2.4), I(2.2), I(7.0), I(7.0),
      fill=None, border=CORAL_LT, bw=Pt(2))
txt(s, I(2.4), I(2.7), I(9.0), I(1.6),
    "One loop. You decide every send.", sz=34, color=WHITE,
    bold=True, font=SERIF, line_spacing=1.15)
rect(s, I(2.45), I(4.45), I(1.2), Pt(2.5), fill=WHITE)
txt(s, I(2.4), I(4.7), I(9), I(0.4),
    "Valori Capital", sz=13, color=WHITE)
add_notes(s, ["Close on the principle."])

# ════════════════════════════════════════════════════════════════════
out = "Valori-fundraising-semi-automatic-pipeline.pptx"
prs.save(out)
print("Saved %d slides to %s" % (len(prs.slides), out))
print("Screenshot placeholders: 4-A, 5-A, 6-A, 7-A, 8-A, 9-A, 9-B, "
      "10-A, 11-A, 12-A, 13-A")
